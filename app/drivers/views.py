import csv
import threading

import openpyxl
from drf_spectacular.utils import extend_schema, extend_schema_view
from rest_framework import status
from rest_framework.exceptions import (
    NotFound,
    PermissionDenied,
    ValidationError,
)
from rest_framework.generics import RetrieveUpdateDestroyAPIView
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from drivers.models import BulkUpload, Driver, DriverStatus, UploadStatus
from drivers.serializers import (
    BulkUploadSerializer,
    BulkUploadStatusSerializer,
    DriverSerializer,
    DriverUpdateSerializer,
)


def get_user_org(user):
    org = getattr(user, 'organization', None)
    if org is None:
        raise PermissionDenied('User not associated with any organization')
    return org


DRIVER_FIELD_MAP = {
    'driver_code': 'driver_code',
    'name': 'name',
    'date_of_birth': 'date_of_birth',
    'license_number': 'license_number',
    'license_class': 'license_class',
    'license_expiry': 'license_expiry',
    'experience_years': 'experience_years',
    'job_id': 'job_id',
    'status': 'status',
    'shift': 'shift',
    'emergency_contact_name': 'emergency_contact_name',
    'emergency_contact_phone': 'emergency_contact_phone',
    'languages': 'languages',
    'user_id': 'user_id',
}


class DriverListCreateView(APIView):
    permission_classes = [IsAuthenticated]

    @extend_schema(
        summary='List drivers',
        description=(
            "Returns all drivers belonging to the "
            "authenticated user's organization."
        ),
        responses={200: DriverSerializer(many=True)},
    )
    def get(self, request):
        org = get_user_org(request.user)
        drivers = Driver.objects.filter(
            organization=org
        ).order_by('-created_at')
        serializer = DriverSerializer(drivers, many=True)
        return Response(serializer.data)

    @extend_schema(
        summary='Create driver',
        description=(
            "Creates a new driver under the authenticated "
            "user's organization."
        ),
        request=DriverSerializer,
        responses={201: DriverSerializer},
    )
    def post(self, request):
        org = get_user_org(request.user)
        serializer = DriverSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        serializer.save(organization=org)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


@extend_schema_view(
    get=extend_schema(
        summary='Get driver by ID',
        description=(
            "Returns a driver by its ID (scoped to "
            "user's organization)."
        ),
        responses={200: DriverSerializer, 404: None},
    ),
    patch=extend_schema(
        summary='Update driver',
        description=(
            "Partially updates a driver (scoped to "
            "user's organization)."
        ),
        request=DriverUpdateSerializer,
        responses={200: DriverSerializer, 404: None},
    ),
    delete=extend_schema(
        summary='Delete driver',
        description=(
            "Deletes a driver (scoped to "
            "user's organization)."
        ),
        responses={204: None, 404: None},
    ),
    put=extend_schema(exclude=True),
)
class DriverDetailView(RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    http_method_names = ['get', 'patch', 'delete']

    def get_serializer_class(self):
        if self.request.method in ('PUT', 'PATCH'):
            return DriverUpdateSerializer
        return DriverSerializer

    def get_queryset(self):
        org = get_user_org(self.request.user)
        return Driver.objects.filter(organization=org)

    def perform_destroy(self, instance):
        instance.delete()

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', True)
        instance = self.get_object()
        serializer = self.get_serializer(
            instance, data=request.data, partial=partial
        )
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        data = DriverSerializer(instance).data
        return Response(data, status=status.HTTP_200_OK)


class BulkUploadView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    @extend_schema(
        summary='Bulk upload drivers',
        description=(
            'Upload a CSV or Excel file containing driver records. '
            'The file is processed asynchronously. Returns a BulkUpload '
            'record with status for polling.'
        ),
        request={
            'multipart/form-data': {
                'type': 'object',
                'properties': {
                    'file': {
                        'type': 'string',
                        'format': 'binary',
                        'description': (
                            'CSV (.csv) or Excel (.xlsx, .xls) file'
                        ),
                    }
                },
                'required': ['file'],
            }
        },
        responses={202: BulkUploadSerializer},
    )
    def post(self, request):
        org = get_user_org(request.user)

        file_obj = request.FILES.get('file')
        if not file_obj:
            raise ValidationError({'file': 'No file provided'})

        ext = (
            file_obj.name.rsplit('.', 1)[-1].lower()
            if '.' in file_obj.name else ''
        )
        if ext not in ('csv', 'xlsx', 'xls'):
            raise ValidationError(
                'Unsupported file format. Use .csv, .xlsx, or .xls'
            )

        bulk = BulkUpload.objects.create(
            organization=org,
            file=file_obj,
        )

        thread = threading.Thread(
            target=_process_bulk_upload,
            args=(bulk.id, ext),
            daemon=True,
        )
        thread.start()

        serializer = BulkUploadSerializer(bulk)
        return Response(serializer.data, status=status.HTTP_202_ACCEPTED)


class BulkUploadStatusView(APIView):
    permission_classes = [IsAuthenticated]

    @extend_schema(
        summary='Get bulk upload status',
        description=(
            'Returns the status, progress, and '
            'error report for a bulk upload.'
        ),
        responses={200: BulkUploadStatusSerializer, 404: None},
    )
    def get(self, request, upload_id):
        org = get_user_org(request.user)
        try:
            bulk = BulkUpload.objects.get(id=upload_id, organization=org)
        except BulkUpload.DoesNotExist:
            raise NotFound('Bulk upload not found')

        serializer = BulkUploadStatusSerializer(bulk)
        return Response(serializer.data)


def _process_bulk_upload(bulk_id, ext):
    from django.db import connection

    try:
        bulk = BulkUpload.objects.get(id=bulk_id)
        bulk.status = UploadStatus.PROCESSING
        bulk.save()

        file_path = bulk.file.path
        rows = []

        if ext == 'csv':
            with open(file_path, newline='', encoding='utf-8-sig') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    rows.append(row)
        elif ext in ('xlsx', 'xls'):
            wb = openpyxl.load_workbook(file_path, read_only=True)
            ws = wb.active
            header_row = next(ws.iter_rows(min_row=1, max_row=1))
            headers = [cell.value for cell in header_row]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append(dict(zip(headers, row)))

        bulk.total_rows = len(rows)
        bulk.save()

        errors = []

        for idx, row in enumerate(rows, start=2):
            try:
                mapped = {}
                for csv_key, model_key in DRIVER_FIELD_MAP.items():
                    value = row.get(csv_key, '')

                    if isinstance(value, str):
                        value = value.strip()

                    if value != '' and value is not None:
                        mapped[model_key] = value

                required_fields = [
                    'driver_code', 'name', 'date_of_birth',
                    'license_number', 'license_class', 'license_expiry',
                    'experience_years',
                ]

                missing = [f for f in required_fields if f not in mapped]
                if missing:
                    raise ValidationError(
                        f'Missing required fields: {", ".join(missing)}'
                    )

                mapped.setdefault('status', DriverStatus.ACTIVE)

                if isinstance(mapped.get('languages'), str):
                    mapped['languages'] = [
                        lang.strip()
                        for lang in mapped['languages'].split(',')
                        if lang.strip()
                    ]

                Driver.objects.create(
                    organization=bulk.organization, **mapped
                )
                bulk.processed_rows += 1

            except Exception as e:
                bulk.failed_rows += 1
                errors.append({
                    'row': idx,
                    'errors': str(e),
                })

        bulk.error_report = errors
        if bulk.failed_rows == 0:
            bulk.status = UploadStatus.COMPLETED
        elif bulk.processed_rows > 0:
            bulk.status = UploadStatus.COMPLETED
        else:
            bulk.status = UploadStatus.FAILED
        bulk.save()

    except Exception as exc:
        try:
            bulk = BulkUpload.objects.get(id=bulk_id)
            bulk.status = UploadStatus.FAILED
            bulk.error_report = bulk.error_report or []
            bulk.error_report.append(
                {'row': 0, 'errors': f'Processing error: {exc}'}
            )
            bulk.save()
        except BulkUpload.DoesNotExist:
            pass
    finally:
        connection.close()
